#!/usr/bin/env python3
"""Export a compact private translation batch with original corpus row IDs."""
import argparse
import csv
import json
import re
import unicodedata
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


def key(text):
    return " ".join(unicodedata.normalize("NFC", text).split())


def export(csv_path, output_dir, limit=100):
    root = Path(__file__).resolve().parents[1]
    known = set()
    for name in ("combined-parallel.jsonl", "opus-en-ks/train-silver.jsonl", "opus-en-ks/evaluation-silver.jsonl"):
        for line in (root / "training-data" / name).read_text().splitlines():
            if line.strip():
                known.add(key(json.loads(line).get("kashmiri", "")))
    records = []
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        for row_id, row in enumerate(csv.DictReader(handle), 1):
            text = " ".join(row["sentence"].split())
            normalized = key(text)
            if "SMUQamar" not in row["source"] or normalized in known:
                continue
            if not 4 <= len(text.split()) <= 12 or len(text) > 150:
                continue
            if re.search(r"[A-Za-z@]|\d", text):
                continue
            known.add(normalized)
            records.append({
                "batch_id": len(records) + 1,
                "corpus_row_id": row_id,
                "kashmiri": text,
                "english": "",
                "source": row["source"],
                "status": "awaiting-user-translation",
            })
            if len(records) == limit:
                break
    if len(records) != limit:
        raise ValueError(f"Only {len(records)} qualifying sentences found")
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"kashmiri-review-batch-{limit}"
    (output_dir / f"{stem}.jsonl").write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in records))
    (output_dir / f"{stem}.txt").write_text("\n\n".join(f"{r['batch_id']}. {r['kashmiri']}" for r in records) + "\n")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Translate"
    sheet.append(["Batch ID", "Corpus row ID", "Kashmiri", "English — fill this", "Source"])
    for record in records:
        sheet.append([record["batch_id"], record["corpus_row_id"], record["kashmiri"], "", record["source"]])
        sheet.cell(record["batch_id"] + 1, 3).alignment = Alignment(horizontal="right", wrap_text=True)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:E{limit + 1}"
    for col, width in (("A", 12), ("B", 18), ("C", 70), ("D", 65), ("E", 65)):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    notes = workbook.create_sheet("Read me")
    notes.append(["Translate by Batch ID 1–100; original CSV row IDs are preserved."])
    notes.append(["Selected from nawabhussain/Kashmiri-Language-Corpus, SMUQamar-attributed rows."])
    notes.append(["Upstream SMUQamar/Kashmiri-English-Parallel-Corpus: CC BY-NC-SA 4.0; retain attribution and terms."])
    notes.append(["Private review batch. Source spellings may need correction. These are not approved training pairs yet."])
    notes.column_dimensions["A"].width = 120
    workbook.save(output_dir / f"{stem}.xlsx")
    print(f"Exported {len(records)} entries; corpus rows {records[0]['corpus_row_id']}–{records[-1]['corpus_row_id']}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("csv", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    export(args.csv, args.output_dir)
