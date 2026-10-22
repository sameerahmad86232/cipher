#!/usr/bin/env python3
import importlib.util,json
from collections import Counter
from openpyxl import load_workbook
from pathlib import Path
root=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('builder',root/'scripts/extract-opus-sentence-candidates.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
rows=m.read(m.OUT/'candidates.jsonl')
parents={(r['source'],r['alignmentLine']):r for r in m.read(root/'training-data/opus-en-ks/train-silver.jsonl')}
protected={m.key(s) for r in m.read(root/'training-data/opus-en-ks/evaluation-silver.jsonl') for s in m.split(r['english'])}
assert len({r['id'] for r in rows})==len(rows)
counts=Counter(m.key(r['english']) for r in rows)
for r in rows:
 p=parents[r['source'],r['alignmentLine']]
 assert r['license']==p['license']
 assert r['parentEnglish']==p['english'] and r['parentKashmiri']==p['kashmiri']
 i=r['segmentIndex']-1
 assert r['english']==m.split(p['english'])[i] and r['kashmiri']==m.split(p['kashmiri'])[i]
 assert m.key(r['english']) not in protected
 assert not r['trainingEligible'] and not r['alignmentVerified']
 assert r['conflictingTargets']==(counts[m.key(r['english'])]>1)
ws=load_workbook(m.OUT/'review.xlsx',read_only=True).active
assert ws.max_row==len(rows)+1
assert all(c.data_type!='f' for row in ws for c in row)
assert len(m.split('One sentence. Another sentence.'))==2
assert len(m.split('اکھ کتھ۔ بیاکھ کتھ۔'))==2
print(f'PASS: {len(rows)} candidate segments match their parents; licenses, evaluation isolation and workbook checked.')
