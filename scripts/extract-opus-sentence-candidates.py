#!/usr/bin/env python3
"""Conservatively segment existing silver paragraphs; outputs require review."""
import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'training-data/opus-sentence-review'
def key(s):return re.sub(r'[^\w\s]','',' '.join(unicodedata.normalize('NFC',s).casefold().split()))
def split(s):return [p.strip() for p in re.split(r'(?<=[.!?۔؟])\s+',s) if p.strip()]
def read(path):return [json.loads(s) for s in path.read_text(encoding='utf-8').splitlines() if s.strip()]
def build():
 OUT.mkdir(parents=True,exist_ok=True)
 blocked=set()
 for p in (ROOT/'training-data').rglob('*.jsonl'):
  if OUT in p.parents:continue
  for r in read(p):
   if isinstance(r,dict) and isinstance(r.get('english'),str):blocked.add(key(r['english']))
 # Exclude evaluation paragraphs AND all their English sentence segments.
 evalrows=read(ROOT/'training-data/opus-en-ks/evaluation-silver.jsonl')
 for r in evalrows:blocked.update(key(s) for s in split(r['english']))
 held=json.loads((ROOT/'training-data/user-reviewed-heldout-grammar.json').read_text())
 blocked.update(key(r['english']) for r in held['records'])
 parents=read(ROOT/'training-data/opus-en-ks/train-silver.jsonl')
 rows=[];seen=set();stats=Counter()
 for r in parents:
  ens,kss=split(r['english']),split(r['kashmiri'])
  if not 2<=len(ens)==len(kss)<=8:continue
  for i,(en,ks) in enumerate(zip(ens,kss),1):
   k=key(en)
   if k in blocked:stats['existingOrEvaluationExcluded']+=1;continue
   if not 3<=len(en.split())<=100 or not 3<=len(ks.split())<=100:continue
   if not .3<=len(ks)/max(1,len(en))<=3:continue
   if any(t in en+ks for t in ('{{','}}','<','>','$','%','http://','https://')):continue
   pair=(k,unicodedata.normalize('NFC',ks))
   if pair in seen:continue
   seen.add(pair)
   rows.append({'id':'opus-sentence-'+hashlib.sha256((en+'\0'+ks).encode()).hexdigest()[:20],'english':en,'kashmiri':ks,'source':r['source'],'sourceVersion':r['sourceVersion'],'license':r['license'],'parentSourceFile':'training-data/opus-en-ks/train-silver.jsonl','alignmentLine':r['alignmentLine'],'segmentIndex':i,'parentEnglish':r['english'],'parentKashmiri':r['kashmiri'],'alignmentMethod':'equal-count-punctuation-segments','review':'pending-native-speaker-review','trainingEligible':False,'alignmentVerified':False})
 targets=Counter(key(r['english']) for r in rows)
 for r in rows:r['conflictingTargets']=targets[key(r['english'])]>1
 (OUT/'candidates.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in rows),encoding='utf-8')
 manifest={'additionalSentenceCandidates':len(rows),'fromExistingSilverParagraphs':True,'newIndependentSourcePairs':False,'sources':dict(Counter(r['source'] for r in rows)),'filters':dict(stats),'automaticTrainingImport':False,'trainingPerformed':False,'warning':'Equal sentence counts do not prove semantic alignment. Abbreviations and punctuation can split incorrectly. Review both parent paragraphs.'}
 (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
 wb=Workbook();ws=wb.active;ws.title='Sentence review'
 columns=['id','english','kashmiri','correction','decision','source','license','alignmentLine','segmentIndex','parentEnglish','parentKashmiri']
 ws.append(columns)
 for r in rows:
  ws.append([r.get(c,'') for c in columns])
  for cell in ws[ws.max_row]:
   if isinstance(cell.value,str):cell.data_type='s'
  ws.cell(ws.max_row,3).alignment=Alignment(readingOrder=2,wrap_text=True)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ('B','C','D','J','K'):ws.column_dimensions[c].width=65
 wb.save(OUT/'review.xlsx')
 print(json.dumps(manifest,indent=2))
if __name__=='__main__':build()
