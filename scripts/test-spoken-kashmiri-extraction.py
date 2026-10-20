#!/usr/bin/env python3
"""Validate a completed private extraction without bundling source book data."""
import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def read_rows(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def validate(output):
    report = json.loads((output / "extraction-report.json").read_text())
    pages = read_rows(output / "pages.jsonl")
    lessons = json.loads((output / "lesson-index.json").read_text())
    assert [page["pdfPage"] for page in pages] == list(range(1, 135))
    assert len(list((output / "fresh-ocr-pages").glob("*.txt"))) == 134
    assert len(list((output / "fresh-ocr-pages").glob("*.tsv"))) == 134
    assert all(page["embeddedText"].strip() and page["archiveOCR"].strip() and page["freshOCR"].strip() for page in pages)
    assert all(page["sourcePage"].endswith(f"/page/n{page['pdfPage'] - 1}/mode/1up") for page in pages)
    assert [lesson["lesson"] for lesson in lessons] == list(range(1, 21))
    assert report["freshOCRPages"] == 134 and not report["trainingEligible"]
    assert "This is a table." in pages[11]["embeddedText"]
    assert "What is this?" in pages[11]["freshOCR"]
    mapping = {
        "vocabulary-candidates": ("vocabularyCandidates", "Vocabulary"),
        "parallel-layout-candidates": ("parallelLayoutCandidates", "Parallel candidates"),
        "grammar-notes": ("grammarNotePageSegments", "Grammar notes"),
        "headings": ("headingCandidates", "Headings"),
        "numeral-candidates": ("numeralCandidates", "Numbers"),
    }
    workbook = load_workbook(output / "spoken-kashmiri-review.xlsx", read_only=True)
    assert workbook["Pages"].max_row == 135
    for filename, (report_key, sheet_name) in mapping.items():
        rows = read_rows(output / f"{filename}.jsonl")
        assert len(rows) == report[report_key] == workbook[sheet_name].max_row - 1
        assert len({row["id"] for row in rows}) == len(rows)
        assert all(1 <= row["pdfPage"] <= 134 and not row["trainingEligible"] for row in rows)
        if filename == "grammar-notes":
            assert {row["lesson"] for row in rows} == set(range(1, 21))
    assert report["cardinalTableLabelChecks"]["missingLabels1to100"] == [22]
    assert report["cardinalTableLabelChecks"]["duplicateLabels"] == [21]
    print("Verified all 134 pages, 20 lessons, OCR files, workbook counts, candidate IDs and source-label anomaly.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    validate(parser.parse_args().output)
