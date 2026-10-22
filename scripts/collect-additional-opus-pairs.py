#!/usr/bin/env python3
"""Collect additional public OPUS alignments for review, never automatic training."""
import hashlib
import json
import re
import unicodedata
import urllib.request
import zipfile
from collections import Counter
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'training-data/additional-opus-review'
CACHE = Path('/workspace/library-files/opus-additional')
SOURCES = {
 'wikimedia': ('v20260327', 'CC BY-SA 4.0'),
 'translatewiki': ('v2026-07-01', 'CC BY 3.0'),
}
def key(s):
    return re.sub(r'[^\w\s]', '', unicodedata.normalize('NFC', s)).casefold().strip()
def build():
    OUT.mkdir(parents=True, exist_ok=True)
    CACHE.mkdir(parents=True, exist_ok=True)
    blocked=set()
    for path in (ROOT/'training-data').rglob('*.jsonl'):
        if OUT in path.parents: continue
        for line in path.read_text(encoding='utf-8').splitlines():
            if not line.strip():continue
            r=json.loads(line)
            if isinstance(r,dict) and isinstance(r.get('english'),str):blocked.add(key(r['english']))
    heldout=json.loads((ROOT/'training-data/user-reviewed-heldout-grammar.json').read_text())
    blocked.update(key(r['english']) for r in heldout['records'])
    rows=[]; seen=set(); stats={}; conflicts=Counter()
    for name,(version,license_) in SOURCES.items():
        url=f'https://object.pouta.csc.fi/OPUS-{name}/{version}/moses/en-ks.txt.zip'
        dest=CACHE/f'{name}-{version}.zip'
        if not dest.exists(): urllib.request.urlretrieve(url,dest)
        with zipfile.ZipFile(dest) as z:
            en_name=next(n for n in z.namelist() if n.endswith('.en'))
            ks_name=next(n for n in z.namelist() if n.endswith('.ks'))
            ens=z.read(en_name).decode('utf-8').splitlines(); kss=z.read(ks_name).decode('utf-8').splitlines()
        assert len(ens)==len(kss)
        counts=Counter()
        for i,(en,ks) in enumerate(zip(ens,kss),1):
            en=' '.join(en.split()); ks=' '.join(ks.split()); k=key(en)
            if k in blocked: counts['existingEnglishExcluded']+=1;continue
            if not 3<=len(en.split())<=120 or not 3<=len(ks.split())<=120:continue
            letters=[c for c in ks if c.isalpha()]
            arabic=sum('\u0600'<=c<='\u06ff' for c in letters)/max(1,len(letters))
            if arabic<.8 or not .22<=len(ks)/max(1,len(en))<=5:continue
            if any(t in en+ks for t in ('{{','}}','<','>','$','%','http://','https://')):continue
            if len(set(ks.split()))/len(ks.split())<.3:continue
            pair=(k,unicodedata.normalize('NFC',ks))
            if pair in seen:continue
            seen.add(pair); conflicts[k]+=1
            rows.append({'id':'opus-review-'+hashlib.sha256((en+'\0'+ks).encode()).hexdigest()[:20], 'english':en,'kashmiri':ks,'source':name,'sourceVersion':version,'sourceURL':url,'alignmentLine':i,'license':license_,'status':'pending-native-speaker-review','languageVerified':False,'trainingEligible':False,'reason':'Additional source alignment outside existing corpus; script filter alone cannot confirm Kashmiri or semantic equivalence.'})
            counts['newReviewCandidates']+=1
        stats[name]={'inputAlignments':len(ens),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),**counts}
    for r in rows:r['conflictingTargets']=conflicts[key(r['english'])]>1
    (OUT/'candidates.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in rows),encoding='utf-8')
    manifest={'additionalPairsForReview':len(rows),'sourceStats':stats,'trainingPerformed':False,'automaticTrainingImport':False,'auditDate':'2026-10-22','unavailableSource':{'repo':'SMUQamar/Kashmiri-English-Parallel-Corpus','gated':'manual','listedLicense':'cc-by-nc-sa-4.0','rawREADMEHTTPStatus':401},'policy':'Retain attribution and share-alike terms; review all new alignments. No gold status or automatic commercial-use clearance.'}
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    wb=Workbook(); ws=wb.active; ws.title='Review candidates'
    ws.append(['ID','English','Kashmiri','Corrected Kashmiri','Review decision','Source','License','Alignment line','Source URL'])
    for r in rows:
        ws.append([r['id'],r['english'],r['kashmiri'],'','pending',r['source'],r['license'],r['alignmentLine'],r['sourceURL']])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
        ws.cell(ws.max_row,3).alignment=Alignment(readingOrder=2,wrap_text=True)
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for column,width in [('A',32),('B',75),('C',75),('D',75),('E',20),('F',18),('G',18),('H',18),('I',55)]:ws.column_dimensions[column].width=width
    wb.save(OUT/'review.xlsx')
    print(json.dumps(manifest,indent=2))
if __name__=='__main__':build()
