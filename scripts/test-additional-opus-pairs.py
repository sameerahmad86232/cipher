#!/usr/bin/env python3
import importlib.util,json
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('collector',root/'scripts/collect-additional-opus-pairs.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
rows=[json.loads(s) for s in (m.OUT/'candidates.jsonl').read_text().splitlines()]
manifest=json.loads((m.OUT/'manifest.json').read_text())
assert rows and len(rows)==manifest['additionalPairsForReview']
assert len({r['id'] for r in rows})==len(rows)
existing=set()
for path in (root/'training-data').rglob('*.jsonl'):
 if m.OUT in path.parents:continue
 for line in path.read_text().splitlines():
  if not line.strip():continue
  r=json.loads(line)
  if isinstance(r,dict) and isinstance(r.get('english'),str):existing.add(m.key(r['english']))
for r in rows:
 assert m.key(r['english']) not in existing
 assert not r['trainingEligible'] and not r['languageVerified']
 assert r['alignmentLine']>0 and r['sourceURL'].startswith('https://object.pouta.csc.fi/OPUS-')
 assert r['license'] in ('CC BY-SA 4.0','CC BY 3.0')
ws=load_workbook(m.OUT/'review.xlsx',read_only=True).active
assert ws.max_row==len(rows)+1
assert all(c.data_type!='f' for row in ws for c in row)
print(f'PASS: {len(rows)} additional review candidates; no existing English overlap, eligibility and spreadsheet checks passed.')
