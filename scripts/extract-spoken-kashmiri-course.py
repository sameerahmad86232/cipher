#!/usr/bin/env python3
"""Page-preserving extraction and independent OCR of IA dli.language.0846.

Requires PyMuPDF, openpyxl and the Tesseract executable with English traineddata.
Outputs are kept in an explicit directory, outside the website/training corpus.
"""
import argparse
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, as_completed
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import xml.etree.ElementTree as ET

import fitz
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

ITEM = "dli.language.0846"
SOURCE = f"https://archive.org/details/{ITEM}"
LESSON_STARTS = [11, 18, 23, 28, 33, 37, 42, 47, 53, 59, 65, 71, 78, 83, 88, 94, 99, 103, 109, 114]
_PDF = None


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def save_jsonl(path, rows):
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def page_reference(pdf_page):
    return f"{SOURCE}/page/n{pdf_page - 1}/mode/1up"


def init_ocr(pdf_path):
    global _PDF
    _PDF = fitz.open(pdf_path)


def ocr_page(arguments):
    index, output_dir, dpi = arguments
    output_dir = Path(output_dir)
    stem = output_dir / f"page-{index + 1:03}"
    image = stem.with_suffix(".png")
    if not stem.with_suffix(".txt").exists() or not stem.with_suffix(".tsv").exists():
        _PDF[index].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY).save(image)
        environment = dict(os.environ, OMP_THREAD_LIMIT="1")
        try:
            subprocess.run(["tesseract", str(image), str(stem), "-l", "eng", "--psm", "3", "txt", "tsv"],
                           check=True, capture_output=True, text=True, env=environment, timeout=180)
        finally:
            image.unlink(missing_ok=True)
    confidences = []
    with stem.with_suffix(".tsv").open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle, delimiter="\t", quoting=csv.QUOTE_NONE):
            if row.get("text", "").strip() and float(row["conf"]) >= 0:
                confidences.append(float(row["conf"]))
    return {"pdfPage": index + 1, "text": stem.with_suffix(".txt").read_text(encoding="utf-8").strip(),
            "meanWordConfidence": round(sum(confidences) / max(len(confidences), 1), 2),
            "recognizedWords": len(confidences)}


def native_lines(page):
    rows = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            text = "".join(span["text"] for span in line["spans"]).strip()
            if text:
                rows.append({"text": text, "bbox": [round(x, 2) for x in line["bbox"]],
                             "fonts": sorted({span["font"] for span in line["spans"]}),
                             "fontSizes": sorted({round(span["size"], 2) for span in line["spans"]})})
    return sorted(rows, key=lambda row: (round(row["bbox"][1] / 3), row["bbox"][0]))


def horizontal_rows(lines):
    result = []
    for line in lines:
        if not 55 < line["bbox"][1] < 735:
            continue
        center = (line["bbox"][1] + line["bbox"][3]) / 2
        near = next((row for row in reversed(result[-3:]) if abs(row["center"] - center) < 3), None)
        if near is None:
            result.append({"center": center, "lines": [line]})
        else:
            near["lines"].append(line)
    for row in result:
        row["lines"].sort(key=lambda line: line["bbox"][0])
        row["text"] = " ".join(line["text"] for line in row["lines"])
    return result


def lesson_number(printed_page):
    if printed_page < 11 or printed_page >= 118:
        return None
    return max(number for number, start in enumerate(LESSON_STARTS, 1) if start <= printed_page)


def candidate(text, pdf_page, kind, section, **fields):
    return {"id": f"{ITEM}-{kind}-{pdf_page}-" + hashlib.sha256(text.encode()).hexdigest()[:12],
            "pdfPage": pdf_page, "printedPage": pdf_page - 1,
            "lesson": lesson_number(pdf_page - 1), "section": section,
            "sourcePage": page_reference(pdf_page), "rawText": text,
            "script": "Roman transcription", "status": "machine-extracted-needs-page-review",
            "trainingEligible": False, **fields}


def extract_candidates(pages):
    headings, vocabulary, parallels, numerals, grammar = [], [], [], [], []
    section = "front matter"
    grammar_section = ""
    for page in pages:
        number = page["pdfPage"]
        printed = number - 1
        appendix = 118 <= printed <= 132
        if printed in LESSON_STARTS:
            section = "dialogue"
            grammar_section = ""
        if appendix:
            section = "appendix"
        rows = horizontal_rows(page["nativeLines"])
        right_starts = Counter(round(line["bbox"][0]) for line in page["nativeLines"]
                               if 280 <= line["bbox"][0] <= 330 and line["bbox"][1] > 120
                               and not line["text"].startswith("Lesson"))
        dialogue_split = right_starts.most_common(1)[0][0] if right_starts else 306
        grammar_rows = []
        for row in rows:
            text = row["text"]
            exact_heading = text.strip(" .:").lower()
            if exact_heading == "vocubulary":
                exact_heading = "vocabulary"
            if exact_heading in ("drills", "exercises", "notes", "vocabulary"):
                section = exact_heading
            is_heading = any(any("Bold" in font for font in line["fonts"]) for line in row["lines"])
            is_subheading = any(any("Italic" in font for font in line["fonts"]) for line in row["lines"])
            if is_heading or (is_subheading and (section == "notes" or appendix)):
                headings.append(candidate(text, number, "heading", section))
                if section == "notes":
                    grammar_section = text
            if section == "notes":
                grammar_rows.append(text)
            if section == "dialogue" and lesson_number(printed):
                left = [line for line in row["lines"] if line["bbox"][0] < dialogue_split - 3]
                right = [line for line in row["lines"] if line["bbox"][0] >= dialogue_split - 3]
                if left and right and not is_heading:
                    roman = re.sub(r"^(?:\d+|[AB])[.)]?\s+", "", " ".join(line["text"] for line in left))
                    english = " ".join(line["text"] for line in right)
                    parallels.append(candidate(text, number, "parallel", section,
                                               romanizedKashmiri=roman, english=english,
                                               alignmentUnit="same-baseline layout candidate; may be sentence fragments"))
            if section in ("vocabulary", "appendix"):
                for half in ([line for line in row["lines"] if line["bbox"][0] < 300],
                             [line for line in row["lines"] if line["bbox"][0] >= 300]):
                    value = " ".join(line["text"] for line in half).strip()
                    quoted = re.match(r"^(.*?)\s*[‘\"]([^’\"]+)[’\"]", value)
                    tagged = re.match(r"^(.+?)\s+([mf](?:\.[sp])?\.?)\s+(.+)$", value)
                    if quoted:
                        form = quoted[1].strip()
                        gender = re.search(r"\s+([mf]\.?)$", form)
                        tag = gender[1] if gender else ""
                        if gender:
                            form = form[:gender.start()].strip()
                        meaning = quoted[2].strip()
                    elif tagged and not is_heading:
                        form, tag, meaning = tagged.groups()
                    else:
                        continue
                    if form and meaning:
                        vocabulary.append(candidate(value, number, "vocabulary", section,
                                                    romanizedKashmiri=form, englishMeaning=meaning,
                                                    grammaticalTag=tag))
                if appendix and not is_heading and not is_subheading:
                    label = r"(?:\.(?!\d)|,(?=\s))"
                    for found in re.finditer(r"(?:^|\s)(\d{1,3})" + label + r"\s*(.+?)(?=\s+\d{1,3}" + label + r"|$)", text):
                        numerals.append(candidate(found.group(0).strip(), number, "numeral", section,
                                                  number=int(found[1]), romanizedKashmiri=found[2].strip()))
                    if printed == 129:
                        for index, line in enumerate(row["lines"][:-1]):
                            if re.fullmatch(r"\d[\d,]{2,}", line["text"]):
                                form = row["lines"][index + 1]["text"]
                                numerals.append(candidate(line["text"] + " " + form, number, "numeral", section,
                                                          number=int(line["text"].replace(",", "")), romanizedKashmiri=form))
        if grammar_rows:
            grammar.append(candidate("\n".join(grammar_rows), number, "grammar", "notes",
                                     heading=grammar_section))
    return tuple(list({record["id"]: record for record in records}.values())
                 for records in (headings, vocabulary, parallels, numerals, grammar))


def review_workbook(output, pages, headings, vocabulary, parallels, numerals, grammar):
    workbook = Workbook()
    info = workbook.active
    info.title = "Read me"
    for text in (
        "Spoken Kashmiri: A Language Course — Omkar N. Koul; first published 1987, second edition 2005.",
        "134 PDF pages. Page 1 is an added cover; printed page = PDF page minus 1.",
        "Kashmiri is in the book's Roman phonetic transcription, NOT Perso-Arabic script.",
        "Preserve raw text: embedded custom fonts and English OCR can both corrupt phonetic symbols.",
        "Vocabulary and same-row English/Roman alignments are review candidates, not verified training pairs.",
        "Control characters unsupported by Excel are omitted only from workbook cells; raw JSON and text preserve them.",
        "Archive metadata reports CC BY-NC 4.0. Keep attribution and the source links; exclude from commercial training.",
        SOURCE,
    ):
        info.append([text])
    info.column_dimensions["A"].width = 115
    sections = {
        "Pages": (["PDF page", "Printed page", "Embedded text", "Archive OCR", "Fresh OCR", "Mean OCR confidence", "Source"],
                  [[p["pdfPage"], p["printedPage"], p["embeddedText"], p["archiveOCR"], p["freshOCR"], p["freshOCRMeanConfidence"], p["sourcePage"]] for p in pages]),
        "Vocabulary": (["ID", "PDF page", "Roman Kashmiri candidate", "English gloss candidate", "Grammar tag", "Your corrected Roman", "Your Perso-Arabic form", "Your English gloss", "Source"],
                       [[p["id"], p["pdfPage"], p["romanizedKashmiri"], p["englishMeaning"], p["grammaticalTag"], "", "", "", p["sourcePage"]] for p in vocabulary]),
        "Parallel candidates": (["ID", "PDF page", "Roman line", "English line", "Corrected Roman", "Corrected English", "Source"],
                                [[p["id"], p["pdfPage"], p["romanizedKashmiri"], p["english"], "", "", p["sourcePage"]] for p in parallels]),
        "Grammar notes": (["ID", "PDF page", "Heading", "Extracted notes", "Your correction", "Source"],
                          [[p["id"], p["pdfPage"], p["heading"], p["rawText"], "", p["sourcePage"]] for p in grammar]),
        "Headings": (["PDF page", "Section", "Heading", "Source"],
                     [[p["pdfPage"], p["section"], p["rawText"], p["sourcePage"]] for p in headings]),
        "Numbers": (["Number", "Roman form candidate", "PDF page", "Corrected form", "Source"],
                    [[p["number"], p["romanizedKashmiri"], p["pdfPage"], "", p["sourcePage"]] for p in numerals]),
    }
    for name, (headers, data) in sections.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row_number, values in enumerate(data, start=2):
            # Preserve raw glyph/control data in JSON/text; Excel forbids some
            # embedded-font control characters in cell strings.
            sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in values])
            for cell in sheet[row_number]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = 14 if "page" in str(column[0].value).lower() else 55
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output / "spoken-kashmiri-review.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_dir", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--workers", type=int, default=3)
    args = parser.parse_args()
    output = args.output_dir
    output.mkdir(parents=True, exist_ok=True)
    ocr_dir = output / "fresh-ocr-pages"
    ocr_dir.mkdir(exist_ok=True)
    pdf_path = args.source_dir / "book.pdf"
    document = fitz.open(pdf_path)
    objects = ET.parse(args.source_dir / "archive-ocr.xml").getroot().findall(".//OBJECT")
    if len(objects) != len(document):
        raise ValueError("Archive OCR/PDF page count mismatch")
    fresh = {}
    with ProcessPoolExecutor(max_workers=args.workers, initializer=init_ocr, initargs=(str(pdf_path),)) as pool:
        futures = [pool.submit(ocr_page, (i, str(ocr_dir), args.dpi)) for i in range(len(document))]
        for future in as_completed(futures):
            result = future.result()
            fresh[result["pdfPage"]] = result
            if len(fresh) % 10 == 0 or len(fresh) == len(document):
                print(f"Fresh OCR: {len(fresh)}/{len(document)} pages", flush=True)
    pages = []
    for i, page in enumerate(document):
        archive_text = "\n".join(" ".join("".join(word.itertext()).strip() for word in line.findall("WORD"))
                                 for line in objects[i].findall(".//LINE"))
        pages.append({"pdfPage": i + 1, "printedPage": i if i else None,
                      "sourcePage": page_reference(i + 1), "embeddedText": page.get_text(),
                      "nativeLines": native_lines(page), "archiveOCR": archive_text,
                      "embeddedControlCharacters": len(ILLEGAL_CHARACTERS_RE.findall(page.get_text())),
                      "freshOCR": fresh[i + 1]["text"],
                      "freshOCRMeanConfidence": fresh[i + 1]["meanWordConfidence"]})
    headings, vocabulary, parallels, numerals, grammar = extract_candidates(pages)
    save_jsonl(output / "pages.jsonl", pages)
    for name, rows in (("headings", headings), ("vocabulary-candidates", vocabulary),
                       ("parallel-layout-candidates", parallels), ("numeral-candidates", numerals),
                       ("grammar-notes", grammar)):
        save_jsonl(output / f"{name}.jsonl", rows)
    for field, filename in (("embeddedText", "embedded-text.txt"), ("archiveOCR", "archive-ocr-page-preserving.txt"), ("freshOCR", "fresh-ocr.txt")):
        (output / filename).write_text("\n\n".join(f"===== PDF PAGE {p['pdfPage']} / PRINTED {p['printedPage']} =====\n{p[field]}" for p in pages), encoding="utf-8")
    token_counts = {field: Counter(token for page in pages
                                   for token in re.findall(r"[^\W\d_][\w:’'\-]*", page[field]))
                    for field in ("embeddedText", "archiveOCR", "freshOCR")}
    all_tokens = sorted(set().union(*(set(counts) for counts in token_counts.values())))
    with (output / "token-inventory.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["rawToken", "embeddedCount", "archiveOCRCount", "freshOCRCount"])
        for token in all_tokens:
            writer.writerow([token, *(counts[token] for counts in token_counts.values())])
    lessons = [{"lesson": i, "printedStart": start,
                "printedEnd": (LESSON_STARTS[i] - 1 if i < len(LESSON_STARTS) else 117),
                "pdfStart": start + 1, "sourcePage": page_reference(start + 1)} for i, start in enumerate(LESSON_STARTS, 1)]
    save_json(output / "lesson-index.json", lessons)
    report = {"title": "Spoken Kashmiri: A Language Course", "author": "Omkar N. Koul",
              "firstPublished": 1987, "editionShownInPDF": 2005, "source": SOURCE,
              "archiveReportedLicense": "CC BY-NC 4.0", "script": "Roman phonetic transcription with English",
              "pdfPages": len(pages), "freshOCRPages": len(fresh), "lessons": len(lessons),
              "vocabularyCandidates": len(vocabulary), "parallelLayoutCandidates": len(parallels),
              "numeralCandidates": len(numerals), "grammarNotePageSegments": len(grammar),
              "headingCandidates": len(headings),
              "unclassifiedRawTokens": len(all_tokens),
              "characters": {field: sum(len(p[field]) for p in pages) for field in ("embeddedText", "archiveOCR", "freshOCR")},
              "embeddedControlCharacters": sum(p["embeddedControlCharacters"] for p in pages),
              "freshOCREngine": subprocess.check_output(["tesseract", "--version"], text=True).splitlines()[0],
              "freshOCRSettings": {"language": "eng", "dpi": args.dpi, "psm": 3},
              "pdfSHA256": hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
              "limitations": ["OCR and embedded custom-font mappings can corrupt phonetic symbols.",
                              "English/Roman same-row candidates may be fragments, not sentence-aligned pairs.",
                              "Candidate counts are not counts of unique verified words or sentences.",
                              "No Perso-Arabic transcription or model training was generated."],
              "trainingEligible": False, "commercialTrainingEligible": False}
    labels = Counter(record["number"] for record in numerals if 1 <= record["number"] <= 100)
    report["cardinalTableLabelChecks"] = {
        "missingLabels1to100": sorted(set(range(1, 101)) - set(labels)),
        "duplicateLabels": sorted(number for number, count in labels.items() if count > 1),
        "note": "Printed page 128 repeats label 21 where 22 is expected; preserve source labels and check the scan before correcting."
    }
    review_workbook(output, pages, headings, vocabulary, parallels, numerals, grammar)
    save_json(output / "extraction-report.json", report)
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
